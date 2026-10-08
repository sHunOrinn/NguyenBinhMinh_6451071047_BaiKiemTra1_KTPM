from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DETAIL = "Chi tiết"
SUMMARY = "Tổng quan"
LAST_ROW = 1000          # cong thuc quet toi dong 1000 de them test case sau van duoc tinh

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
BAND_FILL = PatternFill("solid", start_color="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DETAIL_COLUMNS = [
    ("STT", 6), ("Mã TC", 9), ("Nhóm", 14), ("Tên test case", 38), ("Ưu tiên", 11),
    ("Tiền điều kiện", 28), ("Các bước thực hiện", 46), ("Dữ liệu test", 30),
    ("Kết quả mong đợi", 38), ("Kết quả thực tế", 44), ("Trạng thái", 12),
    ("Thời gian (s)", 12), ("Ảnh chụp lỗi", 24), ("Thời điểm chạy", 19),
]
STATUS_COL = "K"
GROUP_COL = "C"
TIME_COL = "L"
ID_COL = "B"


def _font(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def _rng(col):
    return f"'{DETAIL}'!${col}$2:${col}${LAST_ROW}"


def _write_detail(wb, records):
    ws = wb.create_sheet(DETAIL)
    for c, (name, width) in enumerate(DETAIL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = _font(True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    for i, r in enumerate(records, start=1):
        row = i + 1
        values = [i, r["id"], r["group"], r["title"], r["priority"], r["precondition"],
                  r["steps"], r["data"], r["expected"], r["actual"], r["status"],
                  round(r["duration"], 2), "", r["time"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = _font()
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if c in (1, 2, 5, 11, 12, 14) else "left")
        if r["screenshot"]:
            cell = ws.cell(row=row, column=13, value=Path(r["screenshot"]).name)
            cell.hyperlink = Path(r["screenshot"]).resolve().as_uri()
            cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_COLUMNS))}{max(len(records) + 1, 2)}"

    status_range = f"{STATUS_COL}2:{STATUS_COL}{LAST_ROW}"
    for text, bg, fg in (("PASSED", "C6EFCE", "006100"), ("FAILED", "FFC7CE", "9C0006"),
                         ("ERROR", "FFEB9C", "9C5700"), ("SKIPPED", "D9D9D9", "404040")):
        ws.conditional_formatting.add(status_range, CellIsRule(
            operator="equal", formula=[f'"{text}"'],
            fill=PatternFill("solid", start_color=bg, end_color=bg),
            font=Font(name=FONT, bold=True, color=fg)))
    return ws


def _write_summary(wb, records, meta):
    ws = wb.active
    ws.title = SUMMARY
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", (34, 16, 12, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w

    ws["A1"] = "BÁO CÁO KIỂM THỬ TỰ ĐỘNG – CHỨC NĂNG ĐĂNG NHẬP"
    ws["A1"].font = _font(True, "1F4E78", 14)
    ws.merge_cells("A1:G1")

    info = [
        ("Hệ thống kiểm thử", meta["system"]),
        ("URL", meta["url"]),
        ("Thời điểm chạy", meta["started"]),
        ("Trình duyệt", meta["browser"]),
        ("Chế độ chạy", meta["mode"]),
        ("Công cụ", meta["tools"]),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = _font(True)
        ws.cell(row=i, column=2, value=v).font = _font()
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)

    r0 = 10
    ws.cell(row=r0, column=1, value="KẾT QUẢ TỔNG HỢP").font = _font(True, "1F4E78", 12)
    rows = [
        ("Tổng số test case", f"=COUNTA({_rng(ID_COL)})", "0"),
        ("PASSED (đạt)", f'=COUNTIF({_rng(STATUS_COL)},"PASSED")', "0"),
        ("FAILED (sai kết quả mong đợi)", f'=COUNTIF({_rng(STATUS_COL)},"FAILED")', "0"),
        ("ERROR (lỗi khi chạy)", f'=COUNTIF({_rng(STATUS_COL)},"ERROR")', "0"),
        ("SKIPPED (bỏ qua)", f'=COUNTIF({_rng(STATUS_COL)},"SKIPPED")', "0"),
        ("Tỷ lệ đạt (PASSED / số TC đã chạy)",
         f"=IF(B{r0+1}-B{r0+5}=0,0,B{r0+2}/(B{r0+1}-B{r0+5}))", "0.0%"),
        ("Tổng thời gian chạy (giây)", f"=SUM({_rng(TIME_COL)})", "0.0"),
    ]
    for i, (label, formula, fmt) in enumerate(rows, start=r0 + 1):
        a = ws.cell(row=i, column=1, value=label)
        b = ws.cell(row=i, column=2, value=formula)
        a.font, b.font = _font(), _font(True)
        a.border = b.border = BORDER
        b.number_format = fmt
        b.alignment = Alignment(horizontal="center")
        if i % 2 == 0:
            a.fill = b.fill = BAND_FILL

    g0 = r0 + len(rows) + 3
    ws.cell(row=g0 - 1, column=1, value="THEO NHÓM").font = _font(True, "1F4E78", 12)
    for c, h in enumerate(["Nhóm", "Tổng", "PASSED", "FAILED", "ERROR", "SKIPPED"], start=1):
        cell = ws.cell(row=g0, column=c, value=h)
        cell.font = _font(True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
    groups = list(dict.fromkeys(r["group"] for r in records))
    for i, g in enumerate(groups, start=g0 + 1):
        ws.cell(row=i, column=1, value=g)
        ws.cell(row=i, column=2, value=f"=COUNTIF({_rng(GROUP_COL)},A{i})")
        for c, st in zip(range(3, 7), ("PASSED", "FAILED", "ERROR", "SKIPPED")):
            ws.cell(row=i, column=c,
                    value=f'=COUNTIFS({_rng(GROUP_COL)},$A{i},{_rng(STATUS_COL)},"{st}")')
        for c in range(1, 7):
            cell = ws.cell(row=i, column=c)
            cell.font = _font()
            cell.border = BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")

    note_row = g0 + len(groups) + 3
    notes = [
        "Ghi chú:",
        "• FAILED: ứng dụng cho kết quả khác mong đợi (assert sai).",
        "• ERROR: test không chạy được trọn vẹn (không tìm thấy phần tử, Chrome không mở, mạng lỗi...).",
        "• SKIPPED: bỏ qua có chủ đích, thường do chưa khai báo UTC_USER/UTC_PASS trong file .env.",
        "• Các con số ở đây là công thức, tự cập nhật theo cột 'Trạng thái' của sheet 'Chi tiết'.",
    ]
    for i, text in enumerate(notes):
        ws.cell(row=note_row + i, column=1, value=text).font = _font(i == 0)

    pie = PieChart()
    pie.title = "Phân bố kết quả"
    pie.add_data(Reference(ws, min_col=2, min_row=r0 + 2, max_row=r0 + 5), titles_from_data=False)
    pie.set_categories(Reference(ws, min_col=1, min_row=r0 + 2, max_row=r0 + 5))
    pie.height, pie.width = 7.5, 11
    ws.add_chart(pie, "I3")


def generate_report(records, meta, output_path):
    """records: list[dict]; meta: dict; output_path: duong dan .xlsx. Tra ve Path da ghi."""
    records = sorted(records, key=lambda r: r["id"])
    wb = Workbook()
    _write_summary(wb, records, meta)
    _write_detail(wb, records)
    wb.move_sheet(SUMMARY, offset=-1) if wb.sheetnames[0] != SUMMARY else None
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
